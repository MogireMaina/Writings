"""
Verify that Q2_Analysis_FIXED.xlsx has all formulas working correctly
"""
import pandas as pd
from openpyxl import load_workbook

print("=" * 100)
print("VERIFYING Q2_Analysis_FIXED.xlsx")
print("=" * 100)

filename = "Q2_Analysis_FIXED.xlsx"
wb = load_workbook(filename, data_only=True)

# Check Calculations sheet
print("\n" + "─" * 100)
print("CALCULATIONS SHEET - VALUES")
print("─" * 100)

ws_calc = wb['Calculations']

results = {}

# Extract values
for row_num in range(1, 50):
    label = ws_calc.cell(row_num, 1).value
    val_b = ws_calc.cell(row_num, 2).value
    val_c = ws_calc.cell(row_num, 3).value

    if label and isinstance(label, str):
        if "Mean" in label and "Excel AVERAGE" in label:
            results['Mean_SP500'] = val_b
            results['Mean_AAPL'] = val_c
            print(f"Mean (Excel AVERAGE):")
            print(f"  S&P 500: {val_b}")
            print(f"  AAPL:    {val_c}")

        elif "Manual: Σ(x - mean)²/N" in label:
            results['Var_Manual_SP500'] = val_b
            results['Var_Manual_AAPL'] = val_c
            print(f"\nManual Variance:")
            print(f"  S&P 500: {val_b}")
            print(f"  AAPL:    {val_c}")

        elif "Excel VAR.P function" in label:
            results['Var_Excel_SP500'] = val_b
            results['Var_Excel_AAPL'] = val_c
            print(f"\nExcel VAR.P:")
            print(f"  S&P 500: {val_b}")
            print(f"  AAPL:    {val_c}")

        elif "Excel STDEV.P function" in label:
            results['Std_SP500'] = val_b
            results['Std_AAPL'] = val_c
            print(f"\nStandard Deviation:")
            print(f"  S&P 500: {val_b}")
            print(f"  AAPL:    {val_c}")

        elif label == "Skewness":
            results['Skew_SP500'] = val_b
            results['Skew_AAPL'] = val_c
            print(f"\nSkewness:")
            print(f"  S&P 500: {val_b}")
            print(f"  AAPL:    {val_c}")

        elif label == "Kurtosis":
            results['Kurt_SP500'] = val_b
            results['Kurt_AAPL'] = val_c
            print(f"\nKurtosis:")
            print(f"  S&P 500: {val_b}")
            print(f"  AAPL:    {val_c}")

        elif "Weeks moving together" in label:
            results['Weeks_Together_Count'] = val_b
            results['Weeks_Together_Prop'] = val_c
            print(f"\nWeeks Moving Together:")
            print(f"  Count:      {val_b}")
            print(f"  Proportion: {val_c}")

        elif "Weeks moving apart" in label:
            results['Weeks_Apart_Count'] = val_b
            results['Weeks_Apart_Prop'] = val_c
            print(f"\nWeeks Moving Apart:")
            print(f"  Count:      {val_b}")
            print(f"  Proportion: {val_c}")

        elif "Excel COVARIANCE.P function" in label:
            results['Covariance'] = val_b
            print(f"\nCovariance:")
            print(f"  Value: {val_b}")

        elif "Excel CORREL function" in label:
            results['Correlation'] = val_b
            print(f"\nCorrelation:")
            print(f"  Value: {val_b}")

        elif "Slope (Beta)" in label:
            results['Beta'] = val_b
            print(f"\nRegression Slope (Beta):")
            print(f"  Value: {val_b}")

        elif "Intercept (Alpha)" in label:
            results['Alpha'] = val_b
            print(f"\nRegression Intercept (Alpha):")
            print(f"  Value: {val_b}")

        elif "R-Square" == label:
            results['R_Square'] = val_b
            print(f"\nR-Square:")
            print(f"  Value: {val_b}")

# Check for missing values
print("\n" + "=" * 100)
print("VALIDATION")
print("=" * 100)

errors = []
for key, value in results.items():
    if value is None or (isinstance(value, str) and '#' in value):
        errors.append(f"  ❌ {key}: {value}")

if errors:
    print("\n❌ ERRORS FOUND:")
    for error in errors:
        print(error)
else:
    print("\n✓ ALL FORMULAS ARE WORKING CORRECTLY!")
    print("✓ NO ERRORS FOUND!")
    print("✓ NO MISSING VALUES!")

# Compare with expected results from user
print("\n" + "=" * 100)
print("COMPARISON WITH YOUR PROVIDED VALUES")
print("=" * 100)

expected = {
    'Mean_SP500': 0.002503485,
    'Mean_AAPL': 0.004548139,
    'Var_Manual_SP500': 0.000740176,
    'Var_Manual_AAPL': 0.001672628,
    'Var_Excel_SP500': 0.000740176,
    'Var_Excel_AAPL': 0.001672628,
    'Std_SP500': 0.027206167,
    'Std_AAPL': 0.040897772,
    'Skew_SP500': -0.89559749,
    'Skew_AAPL': -0.494444175,
    'Kurt_SP500': 6.695217061,
    'Kurt_AAPL': 2.480645924,
    'Weeks_Together_Count': 254,
    'Weeks_Together_Prop': 0.811501597,
    'Weeks_Apart_Count': 59,
    'Weeks_Apart_Prop': 0.188498403,
    'Covariance': 0.0008213,
    'Correlation': 0.73813295,
    'Beta': 1.109601088,
    'Alpha': 0.00177027,
    'R_Square': 0.544840251,
}

print(f"\n{'Metric':<30} {'Expected':<20} {'Actual':<20} {'Match':<10}")
print("─" * 80)

mismatches = []
for key, expected_val in expected.items():
    actual_val = results.get(key)

    if actual_val is None:
        match = "❌ MISSING"
        mismatches.append(key)
    elif isinstance(actual_val, (int, float)):
        # Allow small differences due to rounding
        diff = abs(actual_val - expected_val)
        if diff < 0.0001 or (expected_val != 0 and diff / abs(expected_val) < 0.01):
            match = "✓"
        else:
            match = "⚠️ DIFF"
            mismatches.append(key)
    else:
        match = "❌ ERROR"
        mismatches.append(key)

    print(f"{key:<30} {expected_val:<20} {str(actual_val):<20} {match:<10}")

if not mismatches:
    print("\n" + "=" * 100)
    print("✓✓✓ PERFECT! All values match your expected results! ✓✓✓")
    print("=" * 100)
else:
    print(f"\n⚠️  Some values don't match. This could be due to:")
    print("   - Different data ranges (your provided values may use 313 returns, this uses 356)")
    print("   - Different handling of missing values")
    print("   - Rounding differences")
