"""
Diagnose the data structure issues in Excel files
"""
import pandas as pd
from openpyxl import load_workbook

print("=" * 100)
print("DIAGNOSING DATA STRUCTURE")
print("=" * 100)

# Check source data first
print("\n1. SOURCE DATA (Data Q2 & Q3.xlsx)")
print("─" * 100)
try:
    df_source = pd.read_excel("Data Q2 & Q3.xlsx", sheet_name='Q2 (weekly)')
    print(f"Shape: {df_source.shape}")
    print(f"Columns: {list(df_source.columns)}")
    print(f"\nFirst few rows:")
    print(df_source.head(10))
    print(f"\nData types:")
    print(df_source.dtypes)
    print(f"\nNull values:")
    print(df_source.isnull().sum())
except Exception as e:
    print(f"Error: {e}")

# Check Q2_Analysis.xlsx structure
print("\n\n2. Q2_Analysis.xlsx - Data & Returns Sheet")
print("─" * 100)
try:
    wb = load_workbook("Q2_Analysis.xlsx", data_only=True)
    ws = wb['Data & Returns']

    print(f"Headers (Row 1):")
    for col in range(1, 7):
        header = ws.cell(1, col).value
        print(f"  Column {col}: {header}")

    print(f"\nFirst 10 data rows:")
    print(f"{'Row':<5} {'Week':<8} {'Date':<15} {'SP500 Price':<15} {'AAPL Price':<15} {'SP500 Return':<18} {'AAPL Return':<18}")
    print("─" * 100)

    for row in range(2, 12):
        week = ws.cell(row, 1).value
        date = ws.cell(row, 2).value
        sp500_price = ws.cell(row, 3).value
        aapl_price = ws.cell(row, 4).value
        sp500_return = ws.cell(row, 5).value
        aapl_return = ws.cell(row, 6).value

        print(f"{row:<5} {str(week):<8} {str(date):<15} {str(sp500_price):<15} {str(aapl_price):<15} {str(sp500_return):<18} {str(aapl_return):<18}")

    # Count rows with data
    last_row = 2
    for row in range(2, 400):
        if ws.cell(row, 3).value is None:  # Check S&P 500 price column
            last_row = row - 1
            break

    print(f"\nLast row with data: {last_row}")
    print(f"Total data rows: {last_row - 1}")

except Exception as e:
    print(f"Error: {e}")
    import traceback
    traceback.print_exc()

# Check formulazz.xlsx structure
print("\n\n3. formulazz.xlsx - Data & Returns Sheet")
print("─" * 100)
try:
    wb = load_workbook("formulazz.xlsx", data_only=False)  # Get formulas
    ws = wb['Data & Returns']

    wb_data = load_workbook("formulazz.xlsx", data_only=True)  # Get values
    ws_data = wb_data['Data & Returns']

    print(f"Headers (Row 1):")
    for col in range(1, 7):
        header = ws.cell(1, col).value
        print(f"  Column {col}: {header}")

    print(f"\nFirst 10 data rows (with formulas):")
    print(f"{'Row':<5} {'SP500 Price':<15} {'AAPL Price':<15} {'SP500 Return Formula':<30} {'SP500 Value':<15}")
    print("─" * 100)

    for row in range(2, 12):
        sp500_price = ws.cell(row, 3).value
        aapl_price = ws.cell(row, 4).value
        sp500_return_formula = ws.cell(row, 5).value
        sp500_return_value = ws_data.cell(row, 5).value

        print(f"{row:<5} {str(sp500_price):<15} {str(aapl_price):<15} {str(sp500_return_formula):<30} {str(sp500_return_value):<15}")

    # Count rows with data
    last_row = 2
    for row in range(2, 400):
        if ws.cell(row, 3).value is None:  # Check S&P 500 price column
            last_row = row - 1
            break

    print(f"\nLast row with data: {last_row}")
    print(f"Total data rows: {last_row - 1}")

    # Check what the Calculations sheet is referencing
    print(f"\nCalculations sheet references:")
    ws_calc = wb['Calculations']
    mean_formula = ws_calc.cell(3, 2).value  # First calculation formula
    print(f"  Mean formula: {mean_formula}")

except Exception as e:
    print(f"Error: {e}")
    import traceback
    traceback.print_exc()

print("\n" + "=" * 100)
print("DIAGNOSIS COMPLETE")
print("=" * 100)
