import pandas as pd
import numpy as np
from openpyxl import load_workbook
from scipy import stats
import warnings
warnings.filterwarnings('ignore')

# Load the analysis workbook
analysis_file = "Q2_Analysis.xlsx"

print("=" * 80)
print("STEP 1: CALCULATING LOG RETURNS")
print("=" * 80)

# Read the Data & Returns sheet
df_data = pd.read_excel(analysis_file, sheet_name='Data & Returns')

# Calculate log returns
sp500_prices = df_data['S&P 500 Price'].values
aapl_prices = df_data['AAPL Price'].values

# Log returns = ln(P_t / P_t-1)
sp500_returns = np.log(sp500_prices[1:] / sp500_prices[:-1])
aapl_returns = np.log(aapl_prices[1:] / aapl_prices[:-1])

print(f"Number of price observations: {len(sp500_prices)}")
print(f"Number of returns calculated: {len(sp500_returns)}")
print(f"\nFirst 5 S&P 500 returns: {sp500_returns[:5]}")
print(f"First 5 AAPL returns: {aapl_returns[:5]}")

print("\n" + "=" * 80)
print("STEP 2: CALCULATING ALL STATISTICS")
print("=" * 80)

# (a) Mean
sp500_mean = np.mean(sp500_returns)
aapl_mean = np.mean(aapl_returns)
print(f"\n(a) MEAN WEEKLY LOG RETURN")
print(f"S&P 500: {sp500_mean:.10f}")
print(f"AAPL: {aapl_mean:.10f}")

# (b) Variance
# Manual: Σ(x - mean)²/N (population variance)
sp500_var_manual = np.sum((sp500_returns - sp500_mean)**2) / len(sp500_returns)
aapl_var_manual = np.sum((aapl_returns - aapl_mean)**2) / len(aapl_returns)
# Excel VAR.P function (population variance)
sp500_var_excel = np.var(sp500_returns, ddof=0)
aapl_var_excel = np.var(aapl_returns, ddof=0)
print(f"\n(b) VARIANCE CALCULATION")
print(f"S&P 500 - Manual: {sp500_var_manual:.10f}")
print(f"S&P 500 - Excel VAR.P: {sp500_var_excel:.10f}")
print(f"S&P 500 - Difference: {abs(sp500_var_manual - sp500_var_excel):.15f}")
print(f"AAPL - Manual: {aapl_var_manual:.10f}")
print(f"AAPL - Excel VAR.P: {aapl_var_excel:.10f}")
print(f"AAPL - Difference: {abs(aapl_var_manual - aapl_var_excel):.15f}")

# (c) Standard Deviation
# Manual: SQRT(Variance)
sp500_std_manual = np.sqrt(sp500_var_manual)
aapl_std_manual = np.sqrt(aapl_var_manual)
# Excel STDEV.P function (population std dev)
sp500_std_excel = np.std(sp500_returns, ddof=0)
aapl_std_excel = np.std(aapl_returns, ddof=0)
print(f"\n(c) STANDARD DEVIATION CALCULATION")
print(f"S&P 500 - Manual: {sp500_std_manual:.10f}")
print(f"S&P 500 - Excel STDEV.P: {sp500_std_excel:.10f}")
print(f"S&P 500 - Difference: {abs(sp500_std_manual - sp500_std_excel):.15f}")
print(f"AAPL - Manual: {aapl_std_manual:.10f}")
print(f"AAPL - Excel STDEV.P: {aapl_std_excel:.10f}")
print(f"AAPL - Difference: {abs(aapl_std_manual - aapl_std_excel):.15f}")

# (d) Histograms - will be calculated separately using FREQUENCY function bins

# (e) Skewness and Kurtosis
sp500_skew = stats.skew(sp500_returns, bias=False)
aapl_skew = stats.skew(aapl_returns, bias=False)
# Excel kurtosis is excess kurtosis (kurtosis - 3)
sp500_kurt = stats.kurtosis(sp500_returns, bias=False)
aapl_kurt = stats.kurtosis(aapl_returns, bias=False)
print(f"\n(e) SKEWNESS AND KURTOSIS")
print(f"S&P 500 - Skewness: {sp500_skew:.10f}")
print(f"S&P 500 - Kurtosis (excess): {sp500_kurt:.10f}")
print(f"AAPL - Skewness: {aapl_skew:.10f}")
print(f"AAPL - Kurtosis (excess): {aapl_kurt:.10f}")

# (f) Weeks moving together/apart
same_sign = np.sum((sp500_returns > 0) == (aapl_returns > 0))
diff_sign = len(sp500_returns) - same_sign
prop_same = same_sign / len(sp500_returns)
prop_diff = diff_sign / len(sp500_returns)
print(f"\n(f) WEEKS MOVING TOGETHER/APART")
print(f"Weeks moving together: {same_sign} ({prop_same:.10f})")
print(f"Weeks moving apart: {diff_sign} ({prop_diff:.10f})")

# (g) Covariance
# Manual: Σ(x-mean_x)(y-mean_y)/N
cov_manual = np.sum((sp500_returns - sp500_mean) * (aapl_returns - aapl_mean)) / len(sp500_returns)
# Excel COVARIANCE.P function (population covariance)
cov_excel = np.cov(sp500_returns, aapl_returns, ddof=0)[0, 1]
print(f"\n(g) COVARIANCE CALCULATION")
print(f"Manual: {cov_manual:.10f}")
print(f"Excel COVARIANCE.P: {cov_excel:.10f}")
print(f"Difference: {abs(cov_manual - cov_excel):.15f}")

# (h) Correlation
# Manual: Cov(x,y)/(StdDev_x * StdDev_y)
corr_manual = cov_manual / (sp500_std_manual * aapl_std_manual)
# Excel CORREL function
corr_excel = np.corrcoef(sp500_returns, aapl_returns)[0, 1]
print(f"\n(h) CORRELATION CALCULATION")
print(f"Manual: {corr_manual:.10f}")
print(f"Excel CORREL: {corr_excel:.10f}")
print(f"Difference: {abs(corr_manual - corr_excel):.15f}")

# (j) Regression (AAPL vs S&P 500)
# AAPL = alpha + beta * S&P500
# Using numpy polyfit for simple linear regression
slope, intercept = np.polyfit(sp500_returns, aapl_returns, 1)
# Calculate R-squared
y_pred = slope * sp500_returns + intercept
ss_res = np.sum((aapl_returns - y_pred)**2)
ss_tot = np.sum((aapl_returns - aapl_mean)**2)
r_squared = 1 - (ss_res / ss_tot)
print(f"\n(j) REGRESSION STATISTICS (AAPL vs S&P 500)")
print(f"Slope (Beta): {slope:.10f}")
print(f"Intercept (Alpha): {intercept:.10f}")
print(f"R-Square: {r_squared:.10f}")

# (d) Calculate histogram frequencies
print("\n" + "=" * 80)
print("STEP 3: CALCULATING HISTOGRAM FREQUENCIES")
print("=" * 80)

# Read the histogram bins from the Excel file
df_hist = pd.read_excel(analysis_file, sheet_name='Histograms')

# S&P 500 Histogram
# Find where S&P 500 histogram starts
sp500_start = None
aapl_start = None
for i, row in df_hist.iterrows():
    if 'S&P 500 RETURNS HISTOGRAM' in str(row[0]):
        sp500_start = i + 2  # Skip the header row
    if 'AAPL RETURNS HISTOGRAM' in str(row[0]):
        aapl_start = i + 2

print(f"\nS&P 500 histogram starts at row: {sp500_start}")
print(f"AAPL histogram starts at row: {aapl_start}")

# Calculate S&P 500 frequencies
sp500_bins = df_hist.iloc[sp500_start:aapl_start-3, 0].values
sp500_bins = [b for b in sp500_bins if not pd.isna(b)]
print(f"\nS&P 500 bins: {len(sp500_bins)} bins")
print(f"First 5 bins: {sp500_bins[:5]}")
print(f"Last 5 bins: {sp500_bins[-5:]}")

# Use numpy histogram with these bins
# Add -inf at the beginning for the first bin
sp500_bins_with_inf = [-np.inf] + list(sp500_bins)
sp500_freq, _ = np.histogram(sp500_returns, bins=sp500_bins_with_inf)
print(f"\nS&P 500 frequencies calculated: {len(sp500_freq)}")
print(f"Total count: {np.sum(sp500_freq)}")
print(f"First 10 frequencies: {sp500_freq[:10]}")

# Calculate AAPL frequencies
aapl_end = aapl_start
while aapl_end < len(df_hist) and not pd.isna(df_hist.iloc[aapl_end, 0]):
    aapl_end += 1

aapl_bins = df_hist.iloc[aapl_start:aapl_end, 0].values
aapl_bins = [b for b in aapl_bins if not pd.isna(b)]
print(f"\nAAPL bins: {len(aapl_bins)} bins")
print(f"First 5 bins: {aapl_bins[:5]}")
print(f"Last 5 bins: {aapl_bins[-5:]}")

# Use numpy histogram with these bins
aapl_bins_with_inf = [-np.inf] + list(aapl_bins)
aapl_freq, _ = np.histogram(aapl_returns, bins=aapl_bins_with_inf)
print(f"\nAAPL frequencies calculated: {len(aapl_freq)}")
print(f"Total count: {np.sum(aapl_freq)}")
print(f"First 10 frequencies: {aapl_freq[:10]}")

print("\n" + "=" * 80)
print("STEP 4: SAVING ALL RESULTS TO EXCEL")
print("=" * 80)

# Now update the Excel file with all calculations
wb = load_workbook(analysis_file)

# Update Data & Returns sheet with log returns
ws_data = wb['Data & Returns']
# Write S&P 500 returns (starting from row 2, column D)
for i, ret in enumerate(sp500_returns):
    ws_data.cell(row=i+2, column=4, value=ret)
# Write AAPL returns (starting from row 2, column E)
for i, ret in enumerate(aapl_returns):
    ws_data.cell(row=i+2, column=5, value=ret)

# Update Calculations sheet
ws_calc = wb['Calculations']

# Find and update each section
# We'll need to find the row numbers for each calculation
for row_idx, row in enumerate(ws_calc.iter_rows(min_row=1, max_row=100, values_only=False), start=1):
    cell_value = str(row[0].value).strip() if row[0].value else ""

    # (a) Mean
    if cell_value == "Mean (Excel AVERAGE)":
        row[1].value = sp500_mean
        row[2].value = aapl_mean

    # (b) Variance
    elif cell_value == "Manual: Σ(x - mean)²/N":
        row[1].value = sp500_var_manual
        row[2].value = aapl_var_manual
    elif cell_value == "Excel VAR.P function":
        row[1].value = sp500_var_excel
        row[2].value = aapl_var_excel
    elif "Variance" in cell_value and "Difference" in cell_value:
        row[1].value = abs(sp500_var_manual - sp500_var_excel)
        row[2].value = abs(aapl_var_manual - aapl_var_excel)

    # (c) Standard Deviation
    elif cell_value == "Manual: SQRT(Variance)":
        row[1].value = sp500_std_manual
        row[2].value = aapl_std_manual
    elif cell_value == "Excel STDEV.P function":
        row[1].value = sp500_std_excel
        row[2].value = aapl_std_excel
    elif row_idx > 1 and "Standard Deviation" in str(ws_calc.cell(row_idx-1, 1).value or "") and cell_value == "Difference":
        row[1].value = abs(sp500_std_manual - sp500_std_excel)
        row[2].value = abs(aapl_std_manual - aapl_std_excel)

    # (e) Skewness and Kurtosis
    elif cell_value == "Skewness":
        row[1].value = sp500_skew
        row[2].value = aapl_skew
    elif cell_value == "Kurtosis":
        row[1].value = sp500_kurt
        row[2].value = aapl_kurt

    # (f) Weeks moving together/apart
    elif "Weeks moving together" in cell_value:
        row[1].value = same_sign
        row[2].value = prop_same
    elif "Weeks moving apart" in cell_value:
        row[1].value = diff_sign
        row[2].value = prop_diff

    # (g) Covariance
    elif "Manual: Σ(x-mean_x)(y-mean_y)/N" in cell_value or "Manual:" in cell_value and "(x-mean_x)" in cell_value:
        row[1].value = cov_manual
    elif "Excel COVARIANCE.P function" in cell_value:
        row[1].value = cov_excel
    elif row_idx > 1 and "Covariance" in str(ws_calc.cell(row_idx-1, 1).value or "") and cell_value == "Difference":
        row[1].value = abs(cov_manual - cov_excel)

    # (h) Correlation
    elif "Manual: Cov" in cell_value or ("Manual:" in cell_value and "StdDev" in cell_value):
        row[1].value = corr_manual
    elif "Excel CORREL function" in cell_value:
        row[1].value = corr_excel
    elif row_idx > 1 and "Correlation" in str(ws_calc.cell(row_idx-1, 1).value or "") and cell_value == "Difference":
        row[1].value = abs(corr_manual - corr_excel)

    # (j) Regression
    elif "Slope (Beta)" in cell_value or cell_value == "Slope (Beta)":
        row[1].value = slope
    elif "Intercept (Alpha)" in cell_value or cell_value == "Intercept (Alpha)":
        row[1].value = intercept
    elif "R-Square" in cell_value:
        row[1].value = r_squared

# Update Histograms sheet
ws_hist = wb['Histograms']

# Find S&P 500 histogram section and update frequencies
for row_idx, row in enumerate(ws_hist.iter_rows(min_row=1, max_row=200, values_only=False), start=1):
    cell_value = str(row[0].value) if row[0].value else ""
    if 'S&P 500 RETURNS HISTOGRAM' in cell_value:
        # Start updating from 2 rows below
        for i, freq in enumerate(sp500_freq):
            ws_hist.cell(row=row_idx+2+i, column=2, value=int(freq))
        break

# Find AAPL histogram section and update frequencies
for row_idx, row in enumerate(ws_hist.iter_rows(min_row=1, max_row=200, values_only=False), start=1):
    cell_value = str(row[0].value) if row[0].value else ""
    if 'AAPL RETURNS HISTOGRAM' in cell_value:
        # Start updating from 2 rows below
        for i, freq in enumerate(aapl_freq):
            ws_hist.cell(row=row_idx+2+i, column=2, value=int(freq))
        break

# Update Summary Results sheet
ws_summary = wb['Summary Results']

# Find and update each row
for row_idx, row in enumerate(ws_summary.iter_rows(min_row=1, max_row=100, values_only=False), start=1):
    cell_value = str(row[0].value).strip() if row[0].value else ""

    # Dataset Information
    if cell_value == "Stock Selected:":
        row[1].value = "Apple Inc. (AAPL)"
    elif cell_value == "Index Selected:":
        row[1].value = "S&P 500"
    elif cell_value == "Data Frequency:":
        row[1].value = "Weekly"
    elif cell_value == "Period:":
        row[1].value = "November 2019 to November 2025 (6 years)"
    elif cell_value == "Number of Observations:":
        row[1].value = f"{len(sp500_returns)} weekly returns"

    # Statistical Summary
    elif cell_value == "Mean Weekly Return":
        row[1].value = sp500_mean
        row[2].value = aapl_mean
    elif cell_value == "Variance":
        row[1].value = sp500_var_manual
        row[2].value = aapl_var_manual
    elif cell_value == "Standard Deviation":
        row[1].value = sp500_std_manual
        row[2].value = aapl_std_manual
    elif cell_value == "Skewness":
        row[1].value = sp500_skew
        row[2].value = aapl_skew
    elif cell_value == "Kurtosis":
        row[1].value = sp500_kurt
        row[2].value = aapl_kurt
    elif cell_value == "Covariance":
        row[1].value = cov_manual
    elif cell_value == "Correlation":
        row[1].value = corr_excel
    elif cell_value == "Weeks Moving Together":
        row[1].value = same_sign
        row[2].value = prop_same
    elif cell_value == "Weeks Moving Apart":
        row[1].value = diff_sign
        row[2].value = prop_diff

    # Regression Results
    elif "Beta (Slope)" in cell_value or cell_value == "Beta (Slope)":
        row[1].value = slope
    elif "Alpha (Intercept)" in cell_value or cell_value == "Alpha (Intercept)":
        row[1].value = intercept
    elif "R-Square" in cell_value:
        row[1].value = r_squared

# Save the workbook
wb.save(analysis_file)
print(f"\n✓ All calculations saved to {analysis_file}")

print("\n" + "=" * 80)
print("COMPLETE ANALYSIS SUMMARY")
print("=" * 80)
print(f"\nDataset: {len(sp500_returns)} weekly returns")
print(f"\n(a) Mean Weekly Log Return:")
print(f"    S&P 500: {sp500_mean:.10f} ({sp500_mean*100:.4f}%)")
print(f"    AAPL: {aapl_mean:.10f} ({aapl_mean*100:.4f}%)")
print(f"\n(b) Variance:")
print(f"    S&P 500: {sp500_var_manual:.10f}")
print(f"    AAPL: {aapl_var_manual:.10f}")
print(f"\n(c) Standard Deviation:")
print(f"    S&P 500: {sp500_std_manual:.10f} ({sp500_std_manual*100:.4f}%)")
print(f"    AAPL: {aapl_std_manual:.10f} ({aapl_std_manual*100:.4f}%)")
print(f"\n(e) Skewness and Kurtosis:")
print(f"    S&P 500 Skewness: {sp500_skew:.10f}")
print(f"    S&P 500 Kurtosis: {sp500_kurt:.10f}")
print(f"    AAPL Skewness: {aapl_skew:.10f}")
print(f"    AAPL Kurtosis: {aapl_kurt:.10f}")
print(f"\n(f) Weeks Moving Together/Apart:")
print(f"    Together: {same_sign} weeks ({prop_same:.2%})")
print(f"    Apart: {diff_sign} weeks ({prop_diff:.2%})")
print(f"\n(g) Covariance:")
print(f"    {cov_manual:.10f}")
print(f"\n(h) Correlation:")
print(f"    {corr_excel:.10f}")
print(f"\n(j) Regression (AAPL vs S&P 500):")
print(f"    Beta: {slope:.10f}")
print(f"    Alpha: {intercept:.10f}")
print(f"    R-Square: {r_squared:.10f} ({r_squared:.2%})")
print("\n" + "=" * 80)
