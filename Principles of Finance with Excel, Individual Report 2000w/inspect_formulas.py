"""
Detailed inspection of Excel files to identify formula issues and missing values
"""
import pandas as pd
import numpy as np
from openpyxl import load_workbook
import warnings
warnings.filterwarnings('ignore')

files_to_check = ["Q2_Analysis.xlsx", "formulazz.xlsx"]

for filename in files_to_check:
    try:
        print("=" * 100)
        print(f"INSPECTING: {filename}")
        print("=" * 100)

        wb = load_workbook(filename, data_only=False)  # Load with formulas
        wb_data = load_workbook(filename, data_only=True)  # Load with values

        # --- CALCULATIONS SHEET ---
        print("\n" + "─" * 100)
        print("CALCULATIONS SHEET - FORMULA INSPECTION")
        print("─" * 100)

        ws_calc = wb['Calculations']
        ws_calc_data = wb_data['Calculations']

        # Check each section
        sections_to_check = {
            "MEAN": (5, 6),
            "VARIANCE - Manual": (9, 10),
            "VARIANCE - Excel VAR.P": (10, 11),
            "VARIANCE - Difference": (11, 12),
            "STD DEV - Manual": (15, 16),
            "STD DEV - Excel STDEV.P": (16, 17),
            "STD DEV - Difference": (17, 18),
            "SKEWNESS": (21, 22),
            "KURTOSIS": (22, 23),
            "WEEKS TOGETHER - Count": (26, 27),
            "WEEKS TOGETHER - Proportion": (26, 27),
            "WEEKS APART - Count": (27, 28),
            "WEEKS APART - Proportion": (27, 28),
            "COVARIANCE - Manual": (31, 32),
            "COVARIANCE - Excel": (32, 33),
            "COVARIANCE - Difference": (33, 34),
            "CORRELATION - Manual": (37, 38),
            "CORRELATION - Excel": (38, 39),
            "CORRELATION - Difference": (39, 40),
            "REGRESSION - Slope": (43, 44),
            "REGRESSION - Intercept": (44, 45),
            "REGRESSION - R-Square": (45, 46),
        }

        print("\nKey Formulas and Values:")
        print(f"{'Row':<5} {'Label':<40} {'Formula':<50} {'Value B':<15} {'Value C':<15}")
        print("─" * 125)

        for row_num in range(1, 50):
            cell_a = ws_calc.cell(row_num, 1)
            cell_b = ws_calc.cell(row_num, 2)
            cell_c = ws_calc.cell(row_num, 3)

            cell_b_data = ws_calc_data.cell(row_num, 2)
            cell_c_data = ws_calc_data.cell(row_num, 3)

            label = str(cell_a.value)[:40] if cell_a.value else ""

            # Only show rows with formulas or important labels
            if cell_b.value and isinstance(cell_b.value, str) and cell_b.value.startswith('='):
                formula_b = str(cell_b.value)[:50]
                value_b = str(cell_b_data.value)[:15] if cell_b_data.value is not None else "MISSING"
                value_c = str(cell_c_data.value)[:15] if cell_c_data.value is not None else "MISSING"

                # Check for errors
                error_marker = ""
                if cell_b_data.value is None or (isinstance(cell_b_data.value, str) and '#' in cell_b_data.value):
                    error_marker = " ❌ ERROR"

                print(f"{row_num:<5} {label:<40} {formula_b:<50} {value_b:<15} {value_c:<15} {error_marker}")

        # --- DATA & RETURNS SHEET ---
        print("\n" + "─" * 100)
        print("DATA & RETURNS SHEET - LOG RETURN FORMULAS")
        print("─" * 100)

        ws_data = wb['Data & Returns']
        ws_data_values = wb_data['Data & Returns']

        # Check first few log return formulas
        print("\nFirst 5 log return formulas:")
        print(f"{'Row':<5} {'S&P 500 Return Formula':<50} {'AAPL Return Formula':<50}")
        print("─" * 105)

        for row_num in range(3, 8):  # First 5 returns (starting row 3)
            cell_sp500 = ws_data.cell(row_num, 5)  # Column E
            cell_aapl = ws_data.cell(row_num, 6)   # Column F

            sp500_formula = str(cell_sp500.value)[:50] if cell_sp500.value else "MISSING"
            aapl_formula = str(cell_aapl.value)[:50] if cell_aapl.value else "MISSING"

            print(f"{row_num:<5} {sp500_formula:<50} {aapl_formula:<50}")

        # Count returns with values
        returns_with_values = 0
        returns_with_errors = 0
        for row_num in range(3, 316):  # Check all return rows
            val_sp500 = ws_data_values.cell(row_num, 5).value
            val_aapl = ws_data_values.cell(row_num, 6).value

            if val_sp500 is not None and val_aapl is not None:
                if isinstance(val_sp500, (int, float)) and isinstance(val_aapl, (int, float)):
                    returns_with_values += 1
                else:
                    returns_with_errors += 1

        print(f"\nReturns Summary:")
        print(f"  Returns with valid values: {returns_with_values}")
        print(f"  Returns with errors: {returns_with_errors}")

        # --- SUMMARY RESULTS SHEET ---
        print("\n" + "─" * 100)
        print("SUMMARY RESULTS SHEET")
        print("─" * 100)

        ws_summary = wb['Summary Results']
        ws_summary_data = wb_data['Summary Results']

        print("\nKey Metrics:")
        print(f"{'Metric':<35} {'Formula/Value':<40} {'Calculated Value':<20}")
        print("─" * 95)

        for row_num in range(1, 35):
            cell_a = ws_summary.cell(row_num, 1)
            cell_b = ws_summary.cell(row_num, 2)
            cell_c = ws_summary.cell(row_num, 3)

            cell_b_data = ws_summary_data.cell(row_num, 2)
            cell_c_data = ws_summary_data.cell(row_num, 3)

            label = str(cell_a.value)[:35] if cell_a.value else ""

            if label and any(keyword in label for keyword in ['Mean', 'Variance', 'Standard', 'Skewness',
                                                                'Kurtosis', 'Covariance', 'Correlation',
                                                                'Beta', 'Alpha', 'R-Square', 'Weeks']):
                formula = str(cell_b.value)[:40] if cell_b.value else ""
                value = str(cell_b_data.value)[:20] if cell_b_data.value is not None else "MISSING"

                error_marker = ""
                if cell_b_data.value is None or (isinstance(cell_b_data.value, str) and '#' in cell_b_data.value):
                    error_marker = " ❌"

                print(f"{label:<35} {formula:<40} {value:<20} {error_marker}")

        print("\n")

    except Exception as e:
        print(f"❌ Error processing {filename}: {e}")
        import traceback
        traceback.print_exc()

print("\n" + "=" * 100)
print("INSPECTION COMPLETE")
print("=" * 100)
