from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Slide 1: Title Slide
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(2))
title_frame = title_box.text_frame
title_frame.text = "Qatar as a Legal Chokepoint:\nEconomic Warfare and International Law"
title_para = title_frame.paragraphs[0]
title_para.font.size = Pt(36)
title_para.font.bold = True
title_para.alignment = PP_ALIGN.CENTER

subtitle_box = slide1.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(9), Inches(1.5))
subtitle_frame = subtitle_box.text_frame
subtitle_frame.text = "LL.M Thesis Presentation\nResearch Methods in Law"
subtitle_para = subtitle_frame.paragraphs[0]
subtitle_para.font.size = Pt(20)
subtitle_para.alignment = PP_ALIGN.CENTER

# Slide 2: Research Overview & Core Claims
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
title2 = slide2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
title2.text_frame.text = "Core Claims"
title2.text_frame.paragraphs[0].font.size = Pt(28)
title2.text_frame.paragraphs[0].font.bold = True

content2 = slide2.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5.5))
tf2 = content2.text_frame
tf2.word_wrap = True

p1 = tf2.paragraphs[0]
p1.text = "Claim 1: Transnational Legal Exposure"
p1.font.size = Pt(16)
p1.font.bold = True
p1.space_after = Pt(8)

p2 = tf2.add_paragraph()
p2.text = "Qatar shares control of the world's largest gas reserve with Iran, creating unique legal exposure under extraterritorial sanctions (OFAC). These sanctions override territorial sovereignty, reducing Qatar's control from substantive to nominal."
p2.font.size = Pt(14)
p2.level = 1
p2.space_after = Pt(12)

p3 = tf2.add_paragraph()
p3.text = "Claim 2: Contradiction to State Sovereignty"
p3.font.size = Pt(16)
p3.font.bold = True
p3.space_after = Pt(8)

p4 = tf2.add_paragraph()
p4.text = "Economic interdependence at critical chokepoints fundamentally contradicts dominant legal doctrine of state sovereignty. Weaponised interdependence violates UNCLOS Articles 56, 77 and Qatar's Law No. (8) of 2004."
p4.font.size = Pt(14)
p4.level = 1
p4.space_after = Pt(12)

p5 = tf2.add_paragraph()
p5.text = "Claim 3: Multi-Forum Legal Strategy Reveals Gaps"
p5.font.size = Pt(16)
p5.font.bold = True
p5.space_after = Pt(8)

p6 = tf2.add_paragraph()
p6.text = "Qatar's simultaneous litigation (ICJ, ICAO, CERD) during 2017-2021 blockade revealed major gaps in protecting chokepoint states from economic warfare."
p6.font.size = Pt(14)
p6.level = 1

# Slide 3: Methodology & Approach
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
title3 = slide3.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
title3.text_frame.text = "Research Methodology"
title3.text_frame.paragraphs[0].font.size = Pt(28)
title3.text_frame.paragraphs[0].font.bold = True

content3 = slide3.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5.5))
tf3 = content3.text_frame
tf3.word_wrap = True

m1 = tf3.paragraphs[0]
m1.text = "1. Doctrinal Legal Analysis"
m1.font.size = Pt(16)
m1.font.bold = True
m1.space_after = Pt(8)

m2 = tf3.add_paragraph()
m2.text = "Examines sovereignty, jurisdiction, and treaty obligations (UNCLOS, Chicago Convention). Reveals critical gap: remedies remain procedural rather than substantive."
m2.font.size = Pt(14)
m2.level = 1
m2.space_after = Pt(12)

m3 = tf3.add_paragraph()
m3.text = "2. Comparative Case Study"
m3.font.size = Pt(16)
m3.font.bold = True
m3.space_after = Pt(8)

m4 = tf3.add_paragraph()
m4.text = "Compares Qatar (2017-2021), Russia (2022-present), Iran (2011-2025), Taiwan (2022-present). Demonstrates international law distinguishes by legal form rather than economic effect—territorial blockades trigger aviation law, but functional blockades lack legal remedy."
m4.font.size = Pt(14)
m4.level = 1
m4.space_after = Pt(12)

m5 = tf3.add_paragraph()
m5.text = "3. Interdisciplinary Integration (McCrudden Framework)"
m5.font.size = Pt(16)
m5.font.bold = True
m5.space_after = Pt(8)

m6 = tf3.add_paragraph()
m6.text = "Applies 'two-way traffic' between law and social sciences to understand why economic power overrides legal validity. Example: Russian oil price cap works through market control (95% insurance concentration), not legal authority."
m6.font.size = Pt(14)
m6.level = 1

# Slide 4: Key Findings & Legal Gaps
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
title4 = slide4.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
title4.text_frame.text = "Key Findings & Legal Gaps"
title4.text_frame.paragraphs[0].font.size = Pt(28)
title4.text_frame.paragraphs[0].font.bold = True

content4 = slide4.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5.5))
tf4 = content4.text_frame
tf4.word_wrap = True

f1 = tf4.paragraphs[0]
f1.text = "ICJ Jurisdictional Victory Without Substantive Protection"
f1.font.size = Pt(16)
f1.font.bold = True
f1.space_after = Pt(8)

f2 = tf4.add_paragraph()
f2.text = "ICJ affirmed ICAO jurisdiction (2020), but case settled without merits determination or compensation (2021). International law provides formal remedy but no practical protection—reactive, not preventative."
f2.font.size = Pt(14)
f2.level = 1
f2.space_after = Pt(12)

f3 = tf4.add_paragraph()
f3.text = "Form vs. Effect Problem"
f3.font.size = Pt(16)
f3.font.bold = True
f3.space_after = Pt(8)

f4 = tf4.add_paragraph()
f4.text = "Current law distinguishes by legal form (blockade/sanctions/export controls) rather than economic effect. Territorial blockade triggers ICAO; insurance blockade and tech export controls with identical coercive impact have no forum. International law fails to recognize functional equivalence."
f4.font.size = Pt(14)
f4.level = 1
f4.space_after = Pt(12)

f5 = tf4.add_paragraph()
f5.text = "Weaponised Interdependence & Erosion of Sovereignty"
f5.font.size = Pt(16)
f5.font.bold = True
f5.space_after = Pt(8)

f6 = tf4.add_paragraph()
f6.text = "Extraterritorial sanctions (OFAC) create artificial legal boundaries that override sovereign equality. Fishman's 'invisible infrastructure' (dollar clearing, SWIFT, maritime insurance) enables economic warfare without military force, revealing sovereignty is nominal when compliance infrastructures privatize enforcement."
f6.font.size = Pt(14)
f6.level = 1

# Slide 5: Lessons Learned & Feedback Implementation
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
title5 = slide5.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
title5.text_frame.text = "Lessons Learned & Feedback Implementation"
title5.text_frame.paragraphs[0].font.size = Pt(28)
title5.text_frame.paragraphs[0].font.bold = True

content5 = slide5.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5.5))
tf5 = content5.text_frame
tf5.word_wrap = True

l1 = tf5.paragraphs[0]
l1.text = "Key Feedback Integrated"
l1.font.size = Pt(16)
l1.font.bold = True
l1.space_after = Pt(8)

l2 = tf5.add_paragraph()
l2.text = "• Enhanced specificity throughout analysis (FA1, FA8)"
l2.font.size = Pt(14)
l2.level = 1

l3 = tf5.add_paragraph()
l3.text = "• Reorganized research questions from broad to specific (FA6)"
l3.font.size = Pt(14)
l3.level = 1

l4 = tf5.add_paragraph()
l4.text = "• Corrected grammar and terminology (FA7, FA14-FA16)"
l4.font.size = Pt(14)
l4.level = 1

l5 = tf5.add_paragraph()
l5.text = "• Fixed citation formatting—single footnote per sentence (FA17)"
l5.font.size = Pt(14)
l5.level = 1
l5.space_after = Pt(12)

l6 = tf5.add_paragraph()
l6.text = "Critical Methodological Refinement (FA22-FA25)"
l6.font.size = Pt(16)
l6.font.bold = True
l6.space_after = Pt(8)

l7 = tf5.add_paragraph()
l7.text = "Strengthened justification for why doctrinal method connects to claims and research questions. Explained relationship between method and weaponised interdependence—doctrinal analysis alone cannot explain why economic power overcomes legal validity, requiring comparative and interdisciplinary approaches."
l7.font.size = Pt(14)
l7.level = 1
l7.space_after = Pt(12)

l8 = tf5.add_paragraph()
l8.text = "Future Direction: Effect-Based Taxonomy"
l8.font.size = Pt(16)
l8.font.bold = True
l8.space_after = Pt(8)

l9 = tf5.add_paragraph()
l9.text = "Develop empirical thresholds (e.g., 80% insurance market control, 2000 nautical mile rerouting) that translate into justiciable doctrine for courts to apply when chokepoint states allege functional strangulation."
l9.font.size = Pt(14)
l9.level = 1

# Save presentation
prs.save('/home/user/Writings/Qatar_Legal_Chokepoint_Presentation.pptx')
print("Presentation created successfully: Qatar_Legal_Chokepoint_Presentation.pptx")
